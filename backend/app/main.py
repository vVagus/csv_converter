import csv
import io
import os
import random
import re
import uuid
import zipfile
import xml.etree.ElementTree as ET
import docx
from dotenv import load_dotenv

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from app.llm_engine import parse_questions_with_gemini

# Load environment variables from .env
load_dotenv()

app = FastAPI(title="LMS CSV Converter Engine")

# CORS Configuration - Allows cross-origin calls in production
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

STORAGE_DIR = os.path.join(os.path.dirname(__file__), "..", "session_storage")
os.makedirs(STORAGE_DIR, exist_ok=True)


# ==========================================
# HELPER FUNCTIONS & CRASH-PROOF XML PARSER
# ==========================================

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n".join(full_text)


def shuffle_options_and_update_answer(q):
    options_map = {
        "option_1": q.option_1,
        "option_2": q.option_2,
        "option_3": q.option_3,
        "option_4": q.option_4,
    }

    correct_text = options_map.get(q.answear, q.option_1)

    choices = [q.option_1, q.option_2, q.option_3, q.option_4]
    random.shuffle(choices)

    q.option_1 = choices[0]
    q.option_2 = choices[1]
    q.option_3 = choices[2]
    q.option_4 = choices[3]

    new_idx = choices.index(correct_text)
    q.answear = f"option_{new_idx + 1}"

    return q


def snap_to_standard_base(detected_max: float) -> int:
    STANDARD_TARGETS = [20, 25, 30, 40, 50, 100]
    closest = min(STANDARD_TARGETS, key=lambda x: abs(x - detected_max))
    if abs(closest - detected_max) <= 2.0:
        return closest
    return int(round(detected_max))


def parse_xlsx_raw_xml_robust(file_bytes: bytes) -> list[dict]:
    def col_to_index(col_str: str) -> int:
        col_str = re.sub(r'[^A-Z]', '', col_str.upper())
        idx = 0
        for char in col_str:
            idx = idx * 26 + (ord(char) - ord('A') + 1)
        return idx - 1

    ns = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rows_data = []

    with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
        shared_strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            tree = ET.fromstring(z.read('xl/sharedStrings.xml'))
            for si in tree.findall('main:si', ns):
                t_texts = [t.text for t in si.findall('.//main:t', ns) if t.text]
                shared_strings.append("".join(t_texts))

        sheet_xml = None
        for name in z.namelist():
            if name.startswith('xl/worksheets/sheet'):
                sheet_xml = z.read(name)
                break

        if not sheet_xml:
            return []

        tree = ET.fromstring(sheet_xml)
        sheet_data = tree.find('main:sheetData', ns)
        if sheet_data is None:
            return []

        raw_grid = []
        for row in sheet_data.findall('main:row', ns):
            row_dict_by_col = {}
            for cell in row.findall('main:c', ns):
                c_ref = cell.attrib.get('r', '')
                c_idx = col_to_index(c_ref) if c_ref else len(row_dict_by_col)

                c_type = cell.attrib.get('t')
                val = ""

                if c_type == 's':
                    v_elem = cell.find('main:v', ns)
                    if v_elem is not None and v_elem.text.isdigit():
                        s_idx = int(v_elem.text)
                        val = shared_strings[s_idx] if s_idx < len(shared_strings) else ""
                elif c_type == 'inlineStr' or cell.find('main:is', ns) is not None:
                    t_elems = cell.findall('.//main:t', ns)
                    val = "".join([t.text for t in t_elems if t.text])
                else:
                    v_elem = cell.find('main:v', ns)
                    val = v_elem.text if v_elem is not None else ""

                row_dict_by_col[c_idx] = str(val or "").strip()

            if row_dict_by_col:
                max_col = max(row_dict_by_col.keys())
                row_cells = [row_dict_by_col.get(i, "") for i in range(max_col + 1)]
                if any(row_cells):
                    raw_grid.append(row_cells)

        if not raw_grid:
            return []

        header_index = 0
        for i, row in enumerate(raw_grid):
            row_str = " ".join(row).lower()
            if any(k in row_str for k in ["student name", "mark", "score", "status", "register"]):
                header_index = i
                break

        headers = [h.strip() for h in raw_grid[header_index]]

        for r in raw_grid[header_index + 1:]:
            row_dict = {}
            for h, cell_val in zip(headers, r):
                if h:
                    row_dict[h] = cell_val
            if any(row_dict.values()):
                rows_data.append(row_dict)

    return rows_data


def get_field_flexible(row: dict, keywords: list[str], default: str = "") -> str:
    for key, val in row.items():
        clean_key = str(key).lower().strip()
        for kw in keywords:
            if kw in clean_key:
                return str(val).strip()
    return default


# ==========================================
# ROUTE 1: QUESTION BANK CONVERTER
# ==========================================

@app.post("/api/v1/convert")
async def convert_questions(
    subject: str = Form(...),
    question_group: str = Form(...),
    mark: float = Form(...),
    file: UploadFile | None = File(None),
    raw_text: str | None = Form(None),
):
    extracted_text = ""

    if file:
        file_bytes = await file.read()
        extracted_text = extract_text_from_docx(file_bytes)
    elif raw_text and raw_text.strip():
        extracted_text = raw_text.strip()
    else:
        raise HTTPException(
            status_code=400,
            detail="Either a .docx file or raw text input is required.",
        )

    if not extracted_text:
        raise HTTPException(
            status_code=400, detail="Could not extract any text from input source."
        )

    parsed = await parse_questions_with_gemini(extracted_text)

    session_id = str(uuid.uuid4())
    session_folder = os.path.join(STORAGE_DIR, session_id)
    os.makedirs(session_folder, exist_ok=True)

    clean_subject = subject.strip().lower().replace(" ", "_")

    class_map = {}
    for q in parsed.questions:
        cls_name = q.detected_class if q.detected_class else "GENERAL"
        if cls_name not in class_map:
            class_map[cls_name] = []

        shuffled_q = shuffle_options_and_update_answer(q)
        class_map[cls_name].append(shuffled_q)

    generated_files = []

    for cls_name, q_list in class_map.items():
        clean_cls = cls_name.strip().lower().replace(" ", "_")
        filename = f"{clean_subject}_{clean_cls}.csv"
        filepath = os.path.join(session_folder, filename)

        with open(filepath, "w", newline="", encoding="utf-8") as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(
                [
                    "question type",
                    "question group",
                    "question level",
                    "question",
                    "mark",
                    "option_1",
                    "option_2",
                    "option_3",
                    "option_4",
                    "answear",
                ]
            )

            full_group_tag = question_group.strip()

            for q in q_list:
                writer.writerow(
                    [
                        "single_choice",
                        full_group_tag,
                        "medium",
                        q.question,
                        mark,
                        q.option_1,
                        q.option_2,
                        q.option_3,
                        q.option_4,
                        q.answear,
                    ]
                )

        generated_files.append(
            {
                "class_name": cls_name,
                "count": len(q_list),
                "filename": filename,
                "download_url": f"/api/v1/download/{session_id}/{filename}",
            }
        )

    return {
        "session_id": session_id,
        "metrics": {
            "total_questions_processed": parsed.metrics.total_questions_processed,
            "total_answer_corrections": parsed.metrics.total_answer_corrections,
        },
        "csv_files": generated_files,
        "audit_summary": parsed.audit_summary,
    }


@app.get("/api/v1/download/{session_id}/{filename}")
async def download_csv(session_id: str, filename: str):
    path = os.path.join(STORAGE_DIR, session_id, filename)
    if os.path.exists(path):
        return FileResponse(path, filename=filename, media_type="text/csv")
    raise HTTPException(status_code=404, detail="Requested file not found.")


# ==========================================
# ROUTE 2: SCORE NORMALIZER & REGULATOR
# ==========================================

@app.post("/api/v1/normalize-scores")
async def normalize_scores(
    file: UploadFile = File(...),
    target_base: int = Form(0),
    pass_threshold: float = Form(40.0),
    deduplicate: bool = Form(True),
):
    contents = await file.read()
    raw_rows = []

    if file.filename.endswith(".xlsx") or file.filename.endswith(".xls"):
        try:
            raw_rows = parse_xlsx_raw_xml_robust(contents)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not parse Excel spreadsheet: {str(e)}")
    else:
        text = contents.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            raw_rows.append({str(k).strip(): str(v).strip() for k, v in row.items() if k})

    original_row_count = len(raw_rows)

    seen_regs = set()
    clean_rows = []
    for r in raw_rows:
        reg = get_field_flexible(r, ["register", "reg_no", "reg no", "id"])
        if deduplicate and reg:
            if reg in seen_regs:
                continue
            seen_regs.add(reg)
        clean_rows.append(r)

    duplicates_removed = original_row_count - len(clean_rows)
    processed_students = []
    status_corrections = 0

    detected_subject = ""
    detected_class = ""

    for idx, row in enumerate(clean_rows):
        student_name = get_field_flexible(row, ["student name", "full name", "name", "student"], f"Student_{idx+1}")
        reg_no = get_field_flexible(row, ["register no", "register", "reg_no", "reg", "id"], f"REG_{idx+1}")
        lms_status = get_field_flexible(row, ["status", "result", "grade"], "N/A")
        mark_str = get_field_flexible(row, ["mark", "score", "total"], "0/25")

        if not detected_subject:
            detected_subject = get_field_flexible(row, ["subject", "course"], "GENERAL SUBJECT")
        if not detected_class:
            detected_class = get_field_flexible(row, ["class", "grade", "level"], "GENERAL CLASS")

        try:
            if "/" in mark_str:
                num_str, den_str = mark_str.split("/")
                raw_obtained = float(num_str.strip())
                raw_max = float(den_str.strip())
            else:
                raw_obtained = float(mark_str.strip())
                raw_max = 25.0
        except Exception:
            continue

        if raw_max == 0:
            continue

        percentage = (raw_obtained / raw_max) * 100.0
        final_target = (
            target_base if target_base > 0 else snap_to_standard_base(raw_max)
        )

        regulated_score = round((percentage / 100.0) * final_target, 1)
        regulated_mark = f"{regulated_score}/{final_target}"
        regulated_status = "Passed" if percentage >= pass_threshold else "Failed"

        if lms_status.strip().lower() != regulated_status.lower():
            status_corrections += 1

        processed_students.append(
            {
                "studentName": student_name,
                "registerNo": reg_no,
                "originalMark": mark_str,
                "originalStatus": lms_status,
                "percentage": round(percentage, 2),
                "regulatedMark": regulated_mark,
                "regulatedStatus": regulated_status,
            }
        )

    processed_students.sort(key=lambda x: x["percentage"], reverse=True)
    for pos, st in enumerate(processed_students, 1):
        st["position"] = pos

    return {
        "metadata": {
            "subject": detected_subject,
            "class_name": detected_class,
        },
        "metrics": {
            "totalRowsProcessed": original_row_count,
            "cleanCount": len(processed_students),
            "duplicatesRemoved": duplicates_removed,
            "statusCorrections": status_corrections,
        },
        "students": processed_students,
    }


# ==========================================
# ROUTE 3: PDF EXPORT BROADSHEET
# ==========================================

@app.post("/api/v1/export-results-pdf")
async def export_results_pdf(data: dict):
    students = data.get("students", [])
    subject = data.get("subject", "EXAM RESULT BROADSHEET").upper()
    class_name = data.get("class_name", "").upper()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Heading1"],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#111827"),
        alignment=1,
    )

    subtitle_text = f"RESULT BROADSHEET — {subject}"
    if class_name:
        subtitle_text += f" ({class_name})"

    subtitle_style = ParagraphStyle(
        "SubtitleStyle",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#4B5563"),
        alignment=1,
        spaceAfter=15,
    )

    cell_style = ParagraphStyle(
        "CellStyle",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1F2937"),
    )

    elements = [
        Paragraph("<b>STARLITE COLLEGE</b>", title_style),
        Paragraph(f"<b>{subtitle_text}</b>", subtitle_style),
    ]

    table_data = [
        [
            "Pos",
            "Student Name",
            "Register No",
            "Original Mark",
            "Raw %",
            "Regulated Mark",
            "Status",
        ]
    ]

    for st in students:
        table_data.append(
            [
                f"#{st.get('position', '-')}",
                Paragraph(st.get("studentName", ""), cell_style),
                Paragraph(st.get("registerNo", ""), cell_style),
                st.get("originalMark", ""),
                f"{st.get('percentage', 0)}%",
                st.get("regulatedMark", ""),
                st.get("regulatedStatus", ""),
            ]
        )

    table = Table(table_data, colWidths=[40, 200, 140, 90, 60, 110, 80])

    t_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
        ]
    )

    for row_idx, st in enumerate(students, start=1):
        if st.get("regulatedStatus") == "Passed":
            t_style.add(
                "TEXTCOLOR",
                (6, row_idx),
                (6, row_idx),
                colors.HexColor("#15803D"),
            )
        else:
            t_style.add(
                "TEXTCOLOR",
                (6, row_idx),
                (6, row_idx),
                colors.HexColor("#B91C1C"),
            )

    table.setStyle(t_style)
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)

    # Smart file naming featuring Subject AND Class Name
    file_parts = ["Results", subject]
    if class_name:
        file_parts.append(class_name)

    raw_filename = "_".join(file_parts)
    clean_filename = re.sub(r'[^a-zA-Z0-9_\-]', '_', raw_filename) + ".pdf"

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename={clean_filename}"
        },
    )